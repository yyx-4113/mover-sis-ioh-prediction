#!/usr/bin/env python3
"""Convert manuscript_v0.1.md -> manuscript_v0.1.docx with embedded figures and tables.

Faithful conversion: title, author/affiliation block, headings, 6 tables, and 5
figures embedded immediately after each Figure-legend paragraph. The internal
'> ⚠️ TODO' author block is omitted (meta-note, not manuscript content).
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "manuscript", "manuscript_v0.1.md")
FIGDIR = os.path.join(ROOT, "results", "figs")
OUT = os.path.join(ROOT, "manuscript", "manuscript_v0.1.docx")

FIG_MAP = {
    1: "fig1_cohort.png",
    2: "fig2_definitions.png",
    3: "fig3_pr_roc.png",
    4: "fig4_calib_dca.png",
    5: "fig5_importance.png",
}


def add_inline(par, text):
    """Render **bold**, *italic* and ¹ superscript inside a paragraph."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        else:
            iparts = re.split(r"(\*.+?\*)", part)
            for ip in iparts:
                if not ip:
                    continue
                if ip.startswith("*") and ip.endswith("*") and len(ip) > 1:
                    run = par.add_run(ip[1:-1])
                    run.italic = True
                else:
                    segs = re.split(r"(¹)", ip)
                    for s in segs:
                        if s == "¹":
                            r = par.add_run(s)
                            r.font.superscript = True
                        elif s:
                            par.add_run(s)


def is_table_row(line):
    return line.strip().startswith("|")


def parse_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_sep_row(cells):
    return all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c != "")


def add_table(doc, block):
    rows = [parse_table_row(l) for l in block]
    header = rows[0]
    body = [r for r in rows[2:] if any(c for c in r)]
    ncol = len(header)
    body = [r + [""] * (ncol - len(r)) for r in body]
    table = doc.add_table(rows=1, cols=ncol)
    table.style = "Table Grid"
    table.alignment = 1  # center
    hdr = table.rows[0].cells
    for i, c in enumerate(header):
        hdr[i].text = ""
        add_inline(hdr[i].paragraphs[0], c)
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
    for r in body:
        cells = table.add_row().cells
        for i, c in enumerate(r):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], c)
    doc.add_paragraph("")


def main():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    i = 0
    n = len(lines)
    figs_embedded = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped in ("---", "***", "___"):
            i += 1
            continue
        if stripped.startswith("> "):  # internal meta blockquote
            i += 1
            continue

        if stripped.startswith("# ") and not stripped.startswith("## "):
            p = doc.add_paragraph()
            p.style = doc.styles["Title"]
            add_inline(p, stripped[2:])
            i += 1
            continue

        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
            i += 1
            continue

        if is_table_row(line):
            block = []
            while i < n and is_table_row(lines[i]):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue

        m = re.match(r"\*\*(Figure (\d+)\..+?)\*\*", stripped)
        if m:
            fig_no = int(m.group(2))
            p = doc.add_paragraph()
            add_inline(p, stripped)
            # insert the figure image under its legend
            fname = FIG_MAP.get(fig_no)
            if fname:
                fpath = os.path.join(FIGDIR, fname)
                if os.path.exists(fpath):
                    doc.add_picture(fpath, width=Cm(15))
                    doc.paragraphs[-1].alignment = 1  # center
                    figs_embedded.append(fig_no)
            i += 1
            continue

        if stripped == "":
            i += 1
            continue

        # regular paragraph, possibly wrapped across lines
        buf = [stripped]
        j = i + 1
        while j < n:
            nxt = lines[j].strip()
            if nxt == "" or nxt.startswith("#") or is_table_row(lines[j]) \
               or nxt.startswith("> ") or nxt.startswith("---"):
                break
            buf.append(nxt)
            j += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(buf))
        i = j

    doc.save(OUT)
    print("Saved: %s" % OUT)
    print("Figures embedded: %s" % figs_embedded)
    print("Total paragraphs: %d" % len(doc.paragraphs))
    print("Total tables: %d" % len(doc.tables))


if __name__ == "__main__":
    main()
